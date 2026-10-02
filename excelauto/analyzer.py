from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from .models import SheetAnalysis, WritableField, WorkbookAnalysis


class WorkbookAnalyzer:
    """Find conservative form-style writable fields in an xlsx workbook."""

    def __init__(self, workbook) -> None:
        self.workbook = workbook

    @classmethod
    def from_bytes(cls, content: bytes) -> "WorkbookAnalyzer":
        return cls(load_workbook(BytesIO(content), data_only=False, keep_vba=False))

    def analyze(self) -> WorkbookAnalysis:
        return WorkbookAnalysis(sheets=[self._analyze_sheet(ws) for ws in self.workbook.worksheets])

    def _analyze_sheet(self, ws) -> SheetAnalysis:
        fields: list[WritableField] = []
        protected: list[str] = []
        seen_cells: set[str] = set()
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                if cell.data_type == "f" or ws.protection.sheet:
                    protected.append(cell.coordinate)
                    continue
                if not isinstance(cell.value, str) or not cell.value.strip():
                    continue
                target = self._adjacent_empty_cell(ws, cell.row, cell.column)
                if target is None or target.coordinate in seen_cells:
                    continue
                seen_cells.add(target.coordinate)
                fields.append(WritableField(ws.title, target.coordinate, cell.value.strip().rstrip("：:"), "form_field"))
        return SheetAnalysis(name=ws.title, writable_fields=fields, protected_cells=protected)

    @staticmethod
    def _adjacent_empty_cell(ws, row: int, column: int):
        for candidate_column in (column + 1, column + 2):
            candidate = ws.cell(row=row, column=candidate_column)
            if isinstance(candidate, MergedCell):
                continue
            if candidate.value is None and candidate.data_type != "f" and not ws.protection.sheet:
                return candidate
        return None
